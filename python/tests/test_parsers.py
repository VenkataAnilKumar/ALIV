from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from docx import Document

from workers.parsers.docx_parser import parse_docx
from workers.parsers.email_parser import parse_email
from workers.parsers.image_parser import parse_image
from workers.parsers.pdf_parser import parse_pdf


def _create_pdf(path: Path, text: str) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 72), text)
    doc.save(path)
    doc.close()


def test_parse_pdf_extracts_text(tmp_path: Path) -> None:
    pdf_path = tmp_path / "sample.pdf"
    _create_pdf(pdf_path, "Invoice #123")

    parsed = parse_pdf(str(pdf_path))

    assert parsed["parser_name"] == "pymupdf"
    assert "Invoice #123" in parsed["normalized_text"]
    assert parsed["parser_confidence"] > 0
    assert parsed["warnings"] == []


def test_parse_pdf_missing_file_raises() -> None:
    with pytest.raises(FileNotFoundError):
        parse_pdf("does-not-exist.pdf")


def test_parse_docx_extracts_paragraphs_and_table(tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Contract Renewal")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Vendor"
    table.rows[0].cells[1].text = "Contoso"
    doc.save(docx_path)

    parsed = parse_docx(str(docx_path))

    assert parsed["parser_name"] == "python-docx"
    assert "Contract Renewal" in parsed["normalized_text"]
    assert "Vendor | Contoso" in parsed["normalized_text"]
    assert parsed["parser_confidence"] > 0


def test_parse_docx_empty_warns(tmp_path: Path) -> None:
    docx_path = tmp_path / "empty.docx"
    Document().save(docx_path)

    parsed = parse_docx(str(docx_path))

    assert parsed["normalized_text"] == ""
    assert "document produced no extractable text" in parsed["warnings"]


def test_parse_email_extracts_headers_and_body(tmp_path: Path) -> None:
    eml_path = tmp_path / "mail.eml"
    eml_path.write_text(
        "Subject: Hello\\n"
        "From: jane@example.com\\n"
        "To: team@example.com\\n"
        "\\n"
        "This is the body.",
        encoding="utf-8",
    )

    parsed = parse_email(str(eml_path))

    assert parsed["parser_name"] == "email-rfc822"
    assert "Subject: Hello" in parsed["normalized_text"]
    assert "This is the body." in parsed["normalized_text"]
    assert parsed["parser_confidence"] > 0


def test_parse_email_empty_warns(tmp_path: Path) -> None:
    eml_path = tmp_path / "empty.eml"
    eml_path.write_text("", encoding="utf-8")

    parsed = parse_email(str(eml_path))

    assert parsed["normalized_text"] == ""
    assert "email produced no extractable text" in parsed["warnings"]


def test_parse_image_uses_ocr_adapter(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    image_path = tmp_path / "sample.png"
    image_path.write_bytes(b"png-bytes")

    monkeypatch.setattr(
        "workers.parsers.image_parser.run_tesseract",
        lambda _: {
            "recognized_text": "scanned content",
            "confidence_map": {"mean_confidence": 82.5, "word_count": 2},
            "quality_flags": [],
        },
    )

    parsed = parse_image(str(image_path))

    assert parsed["parser_name"] == "tesseract-image"
    assert parsed["normalized_text"] == "scanned content"
    assert parsed["parser_confidence"] == 0.825
    assert parsed["warnings"] == []
