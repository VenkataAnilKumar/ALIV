from __future__ import annotations

from pathlib import Path

from docx import Document


def parse_docx(file_path: str) -> dict:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    document = Document(path)
    text_blocks = _collect_text_blocks(document)
    warnings: list[str] = []

    if not text_blocks:
        warnings.append("document produced no extractable text")

    pages: list[dict] = []
    page_map: list[dict] = []
    segments = _chunk_blocks(text_blocks, chunk_size=8)
    for segment_index, segment in enumerate(segments):
        text = "\n\n".join(segment).strip()
        block_count = len(segment)
        page_entry = {
            "page_number": segment_index,
            "text": text,
            "block_count": block_count,
            "has_text": bool(text),
        }
        pages.append(page_entry)
        page_map.append(
            {
                "page_number": segment_index,
                "text_length": len(text),
                "block_count": block_count,
                "has_text": bool(text),
            }
        )

    normalized_text = "\n\n".join(text_blocks)
    parser_confidence = _estimate_parser_confidence(text_blocks, normalized_text)

    return {
        "file_path": str(path),
        "parser_name": "python-docx",
        "normalized_text": normalized_text,
        "page_map": page_map,
        "pages": pages,
        "parser_confidence": parser_confidence,
        "warnings": warnings,
    }


def _collect_text_blocks(document: Document) -> list[str]:
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                blocks.append(row_text)

    return blocks


def _chunk_blocks(blocks: list[str], chunk_size: int) -> list[list[str]]:
    if not blocks:
        return []

    return [blocks[index:index + chunk_size] for index in range(0, len(blocks), chunk_size)]


def _estimate_parser_confidence(text_blocks: list[str], normalized_text: str) -> float:
    if not text_blocks:
        return 0.0

    block_signal = min(len(text_blocks) / 20, 1.0)
    text_signal = min(len(normalized_text) / 5000, 1.0)
    confidence = (0.55 * block_signal) + (0.45 * text_signal)
    return round(confidence, 4)
